"""Build the versioned D&O knowledge base consumed by the API (ADR-017).

Sources, in precedence order:
  1. docs/domain/DO_KNOWLEDGE_BASE.md — catalog (5.1, 5.2) and weight matrix (6.1);
  2. docs/domain/sources/01_matriz_equivalencia_dicionario_do.xlsx — variants,
     model question, extraction rule and validation rule of DO-001..DO-035.

Output: app/infrastructure/knowledge_base/knowledge_base.json

Usage (from backend/):
    python scripts/build_knowledge_base.py --version 2026.09.1
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

BACKEND_DIR = Path(__file__).resolve().parents[1]
REPO_DIR = BACKEND_DIR.parent
KB_MARKDOWN = REPO_DIR / "docs" / "domain" / "DO_KNOWLEDGE_BASE.md"
DICTIONARY_XLSX = (
    REPO_DIR / "docs" / "domain" / "sources" / "01_matriz_equivalencia_dicionario_do.xlsx"
)
OUTPUT = BACKEND_DIR / "app" / "infrastructure" / "knowledge_base" / "knowledge_base.json"

IMPORTANCE_BY_WEIGHT = {10: "CRITICAL", 7: "HIGH", 4: "MEDIUM", 2: "LOW"}
EXPECTED_TOTAL_WEIGHT = 207
EXTRACTION_RULE = "Extrair trecho literal, fonte, cláusula, página e status de contratação."
DEFAULT_CRITERION = "Comparar redação, alcance, gatilhos, limites, exclusões e contratação."

# Evaluation criteria for the concepts added by the weight matrix (knowledge base, 6.5).
NEW_CONCEPT_CRITERIA = {
    "DO-036": "Valor, moeda, limite agregado ou por evento, sublimites, erosão, franquia, "
    "participação obrigatória e relação com o LMG total.",
    "DO-037": "Data ou extensão, se é automática ou contratável, condições de elegibilidade.",
    "DO-038": "Território, local do ato e da reclamação, jurisdição e limitações por sanções.",
    "DO-039": "Extensão, eventos que ativam o prazo, notificação e elegibilidade.",
    "DO-040": "Extensão, contratação mediante prêmio adicional e condições.",
    "DO-041": "Prazo, forma e efeitos da notificação de fatos ou circunstâncias.",
    "DO-042": "Cobertura, exclusões e condições para atos lesivos (Lei 12.846/2013).",
    "DO-043": "Território e jurisdições aceitos para reclamações no exterior.",
    "DO-044": "Despesas elegíveis e condições de salvamento e contenção.",
}


def table_rows(markdown: str, heading: str) -> list[list[str]]:
    """Return the data rows of the first markdown table after `heading`."""

    start = markdown.index(heading)
    rows: list[list[str]] = []
    in_table = False
    for line in markdown[start:].splitlines()[1:]:
        if line.startswith("|"):
            in_table = True
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                rows.append(cells)
        elif in_table:
            break
    return rows[1:]  # drop header


def split_variants(raw: str | None) -> list[str]:
    if not raw:
        return []
    return [part.strip().rstrip(".") for part in str(raw).split(";") if part.strip()]


def dedupe(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        key = value.lower()
        if key and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def read_dictionary() -> dict[str, dict[str, object]]:
    workbook = load_workbook(DICTIONARY_XLSX, read_only=True, data_only=True)
    training = workbook["02 Base de Treinamento"]
    comparative = workbook["01 Dicionário Comparativo"]

    criteria = [
        row[4] for row in comparative.iter_rows(min_row=2, values_only=True) if row and row[1]
    ]
    entries: dict[str, dict[str, object]] = {}
    for index, row in enumerate(training.iter_rows(min_row=2, values_only=True)):
        if not row or not row[0] or not str(row[0]).startswith("DO-"):
            continue
        concept_id, _domain, canonical, variants, question, rule, *_rest = row
        entries[str(concept_id)] = {
            "canonical": str(canonical).rstrip("."),
            "variants": split_variants(variants),
            "model_question": question,
            "extraction_rule": rule,
            "criterion": criteria[index] if index < len(criteria) else None,
        }
    return entries


def build(version: str) -> dict[str, object]:
    markdown = KB_MARKDOWN.read_text(encoding="utf-8")
    dictionary = read_dictionary()

    concepts: dict[str, dict[str, object]] = {}
    for concept_id, domain, name, _status, review in table_rows(markdown, "### 5.1"):
        source = dictionary.get(concept_id, {})
        concepts[concept_id] = {
            "id": concept_id,
            "domain": domain,
            "name": name,
            "variants": [name, str(source.get("canonical", "")), *source.get("variants", [])],  # type: ignore[misc]
            "criterion": source.get("criterion") or DEFAULT_CRITERION,
            "human_review": "REQUIRED" if review.lower() == "sim" else "RECOMMENDED",
            "pending_validation": False,
            "model_question": source.get("model_question"),
            "extraction_rule": source.get("extraction_rule"),
        }
    for concept_id, domain, name, variants in table_rows(markdown, "### 5.2"):
        concepts[concept_id] = {
            "id": concept_id,
            "domain": domain,
            "name": name,
            "variants": [name, *[v.strip() for v in variants.split(";")]],
            "criterion": NEW_CONCEPT_CRITERIA.get(concept_id, DEFAULT_CRITERION),
            "human_review": "REQUIRED",
            "pending_validation": True,
            "model_question": f"Qual é o tratamento de “{name}” no documento?",
            "extraction_rule": EXTRACTION_RULE,
        }

    for concept in concepts.values():
        concept.update(
            {"weight": None, "importance": "UNWEIGHTED", "justification": None, "related": []}
        )

    for weight, _level, weighted_name, ids_cell, justification in table_rows(markdown, "### 6.1"):
        ids = re.findall(r"(DO-\d{3})(?:\s*\((principal|relacionado)\))?", ids_cell)
        principal = next((cid for cid, role in ids if role in ("", "principal")), ids[0][0])
        related = [cid for cid, _role in ids if cid != principal]
        concept = concepts[principal]
        concept.update(
            {
                "weight": int(weight),
                "importance": IMPORTANCE_BY_WEIGHT[int(weight)],
                "justification": justification,
                "related": related,
                "name": weighted_name,
                "variants": [weighted_name, *concept["variants"]],  # type: ignore[list-item]
            }
        )
        for related_id in related:
            concepts[related_id]["related"] = [principal]

    for concept in concepts.values():
        concept["variants"] = dedupe([str(v) for v in concept["variants"] if v])  # type: ignore[union-attr]

    ordered = sorted(concepts.values(), key=lambda concept: str(concept["id"]))
    total = sum(int(c["weight"]) for c in ordered if c["weight"] is not None)  # type: ignore[call-overload]
    if total != EXPECTED_TOTAL_WEIGHT:
        raise SystemExit(f"Soma dos pesos {total} diferente de {EXPECTED_TOTAL_WEIGHT}.")
    return {"version": version, "concepts": ordered}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--version", default="2026.09.1")
    args = parser.parse_args()

    knowledge_base = build(args.version)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        json.dumps(knowledge_base, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    concepts = knowledge_base["concepts"]
    weighted = [c for c in concepts if c["weight"] is not None]  # type: ignore[union-attr, index]
    print(f"{len(concepts)} conceitos, {len(weighted)} com peso -> {OUTPUT.relative_to(REPO_DIR)}")  # type: ignore[arg-type]
    return 0


if __name__ == "__main__":
    sys.exit(main())
