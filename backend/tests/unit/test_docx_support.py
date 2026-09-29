import io
import zipfile
from pathlib import Path

import pytest
from app.application.commands import CreatePolicyCommand, UploadedFile
from app.application.file_types import DOCX_MIME, detect_content_type
from app.application.use_cases import PolicyService, UploadLimits
from app.domain.interfaces.ports import DocumentUnreadableError
from app.domain.value_objects import DocumentStatus, DocumentType, FileKind, PolicyStatus
from app.infrastructure.events import InMemoryEventBus
from app.infrastructure.knowledge_base import JsonConceptCatalog
from app.infrastructure.persistence import InMemoryPolicyRepository
from app.infrastructure.storage import LocalBlobStorage
from app.infrastructure.word import PythonDocxTextReader
from app.shared.exceptions import ApplicationError
from docx import Document

from tests.fakes import PDF_BYTES, FakePdfReader, RecordingExtractor, make_docx


def _zip(files: dict[str, str]) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return stream.getvalue()


def _office_zip(main_part: str, content_type: str) -> bytes:
    types = (
        '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
        f'content-types"><Override PartName="/{main_part}" ContentType="{content_type}"/></Types>'
    )
    return _zip({"[Content_Types].xml": types, main_part: "<x/>"})


XLSX = _office_zip(
    "xl/workbook.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)
PPTX = _office_zip(
    "ppt/presentation.xml",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
)


def _service(tmp_path: Path, extractor: RecordingExtractor) -> PolicyService:
    return PolicyService(
        repository=InMemoryPolicyRepository(),
        storage=LocalBlobStorage(tmp_path),
        catalog=JsonConceptCatalog(),
        extractor=extractor,
        pdf_reader=FakePdfReader(),
        docx_reader=PythonDocxTextReader(),
        event_bus=InMemoryEventBus(),
        limits=UploadLimits(max_file_bytes=1024 * 1024, max_files=5, min_evidence_confidence=0.7),
    )


def _command(*files: tuple[str, bytes]) -> CreatePolicyCommand:
    return CreatePolicyCommand(
        insurer=None,
        name=None,
        correlation_id="corr-1",
        files=[
            UploadedFile(
                filename=name,
                declared_content_type="application/octet-stream",
                data=data,
                document_type=DocumentType.POLICY,
            )
            for name, data in files
        ],
    )


# ---------- detection ----------


def test_valid_docx_is_detected_by_content() -> None:
    assert detect_content_type(make_docx()) == DOCX_MIME


@pytest.mark.parametrize(
    "data",
    [
        _zip({"leia-me.txt": "oi"}),
        XLSX,
        PPTX,
        _zip({"[Content_Types].xml": "<Types/>", "word/document.xml": "<x/>"}),
        b"PK\x03\x04 truncado",
    ],
)
def test_other_zips_are_not_docx(data: bytes) -> None:
    assert detect_content_type(data) is None


# ---------- reader ----------


def test_reader_extracts_structure_in_order_with_logical_pages() -> None:
    result = PythonDocxTextReader().read(make_docx())

    assert result.page_count == 2  # the explicit page break splits the document
    first, second = result.page_texts[1], result.page_texts[2]
    assert first.startswith("[Cabeçalho] Cabeçalho da Apólice")
    assert first.index("# Condições Gerais") < first.index("1.1 A seguradora")
    assert "- Custos de defesa" in first
    assert "Linha 2: [Coluna 1] LMG | [Coluna 2] R$ 10.000.000" in first
    assert "2.1 Exclusões" in second and "Exclusões" not in first


def test_reader_includes_footer_and_splits_long_blocks() -> None:
    document = Document()
    document.sections[0].footer.paragraphs[0].text = "Rodapé 123"
    for index in range(6):
        document.add_paragraph(f"Parágrafo {index} " + "x" * 40)
    stream = io.BytesIO()
    document.save(stream)

    result = PythonDocxTextReader(max_chars_per_page=120).read(stream.getvalue())

    assert result.page_count > 1
    assert "[Rodapé] Rodapé 123" in result.page_texts[result.page_count]
    assert PythonDocxTextReader(max_chars_per_page=120).read(stream.getvalue()) == result


def test_reader_rejects_corrupted_docx() -> None:
    with pytest.raises(DocumentUnreadableError):
        PythonDocxTextReader().read(b"PK\x03\x04 lixo")


# ---------- upload validation ----------


async def test_upload_accepts_docx_and_stores_with_docx_extension(tmp_path: Path) -> None:
    service = _service(tmp_path, RecordingExtractor())

    policy = await service.create_policy(_command(("apolice.docx", make_docx())))

    document = policy.documents[0]
    assert document.content_type == DOCX_MIME
    assert document.file_kind == FileKind.DOCX
    assert document.ocr_required is False
    assert document.storage_key.endswith(".docx")


@pytest.mark.parametrize(
    ("name", "data", "code", "status"),
    [
        ("dados.zip", _zip({"a.txt": "oi"}), "UNSUPPORTED_MEDIA_TYPE", 415),
        ("planilha.xlsx", XLSX, "UNSUPPORTED_MEDIA_TYPE", 415),
        ("slides.pptx", PPTX, "UNSUPPORTED_MEDIA_TYPE", 415),
        ("quebrado.docx", b"PK\x03\x04 truncado", "DOCX_CORRUPTED", 422),
        (
            "protegido.docx",
            b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64,
            "DOCX_PROTECTED",
            422,
        ),
    ],
)
async def test_upload_rejects_unsupported_or_broken_office_files(
    tmp_path: Path, name: str, data: bytes, code: str, status: int
) -> None:
    service = _service(tmp_path, RecordingExtractor())

    with pytest.raises(ApplicationError) as error:
        await service.create_policy(_command((name, data)))

    assert error.value.code == code
    assert error.value.status_code == status
    assert "Formato" in error.value.message or name in error.value.message


async def test_upload_rejects_docx_whose_body_is_not_parseable(tmp_path: Path) -> None:
    valid = zipfile.ZipFile(io.BytesIO(make_docx()))
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as broken:
        for item in valid.infolist():
            content = b"<w:document" if item.filename == "word/document.xml" else valid.read(item)
            broken.writestr(item.filename, content)
    service = _service(tmp_path, RecordingExtractor())

    with pytest.raises(ApplicationError) as error:
        await service.create_policy(_command(("apolice.docx", stream.getvalue())))

    assert error.value.code == "DOCX_CORRUPTED"


async def test_upload_size_limit_still_applies_to_docx(tmp_path: Path) -> None:
    service = _service(tmp_path, RecordingExtractor())
    service._limits = UploadLimits(max_file_bytes=100, max_files=5, min_evidence_confidence=0.7)

    with pytest.raises(ApplicationError) as error:
        await service.create_policy(_command(("apolice.docx", make_docx())))

    assert error.value.code == "FILE_TOO_LARGE"


# ---------- pipeline ----------


async def test_docx_text_follows_native_path_without_ocr(tmp_path: Path) -> None:
    extractor = RecordingExtractor()
    service = _service(tmp_path, extractor)
    policy = await service.create_policy(_command(("apolice.docx", make_docx())))

    await service.process_policy(policy.id)

    saved = await service.get_policy(policy.id)
    document = saved.documents[0]
    assert document.status == DocumentStatus.COMPLETED
    assert document.file_kind == FileKind.DOCX
    assert document.ocr_required is False
    assert document.pages == 2
    assert set(extractor.contents[0].page_texts) == {1, 2}
    assert saved.status != PolicyStatus.FAILED


async def test_pdf_keeps_working_and_mixed_policy_processes(tmp_path: Path) -> None:
    extractor = RecordingExtractor()
    service = _service(tmp_path, extractor)
    policy = await service.create_policy(
        _command(("apolice.pdf", PDF_BYTES), ("condicoes.docx", make_docx()))
    )

    await service.process_policy(policy.id)

    kinds = {d.filename: d.file_kind for d in (await service.get_policy(policy.id)).documents}
    assert kinds["apolice.pdf"] == FileKind.SCANNED_PDF  # FakePdfReader has no text layer
    assert kinds["condicoes.docx"] == FileKind.DOCX


async def test_docx_without_text_fails_the_document_with_clear_message(tmp_path: Path) -> None:
    empty = io.BytesIO()
    Document().save(empty)
    service = _service(tmp_path, RecordingExtractor())
    policy = await service.create_policy(_command(("vazio.docx", empty.getvalue())))

    await service.process_policy(policy.id)

    document = (await service.get_policy(policy.id)).documents[0]
    assert document.status == DocumentStatus.FAILED
    assert document.failure is not None and "não contém texto" in document.failure
