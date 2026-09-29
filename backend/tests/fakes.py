"""Test doubles for the AI ports. They exist only in tests; the app never uses them."""

import io
from pathlib import Path

from app.application.use_cases import (
    COMPARISON_REQUESTED,
    POLICY_UPLOADED,
    ComparisonService,
    ConceptService,
    PolicyService,
    UploadLimits,
)
from app.domain.entities import (
    AssessmentDecision,
    ConceptAssessment,
    ConceptOccurrence,
    DocumentIntake,
    Evidence,
    ExecutiveSummary,
    ExtractionResult,
    KnowledgeBase,
)
from app.domain.interfaces.ports import AssessmentRequest, DocumentContent, PdfText
from app.domain.services.scoring import ScoringParameters
from app.domain.value_objects import (
    ContractStatus,
    DocumentType,
    Level,
    OccurrenceType,
    TermRelation,
)
from app.infrastructure.events import QueuedEventBus
from app.infrastructure.knowledge_base import JsonConceptCatalog
from app.infrastructure.persistence import InMemoryComparisonRepository, InMemoryPolicyRepository
from app.infrastructure.storage import LocalBlobStorage
from app.infrastructure.word import PythonDocxTextReader
from app.presentation.api.app import create_app
from app.presentation.api.dependencies import ApiServices
from app.shared.config.settings import Settings
from docx import Document
from fastapi import FastAPI


def occurrence(
    concept_id: str,
    status: ContractStatus,
    document_id: str = "doc",
    *,
    amount: str | None = None,
    limit_basis: str | None = None,
) -> ConceptOccurrence:
    evidence = (
        []
        if status == ContractStatus.NOT_FOUND
        else [
            Evidence(
                id=f"{document_id}_{concept_id}",
                document_id=document_id,
                document_name=f"{document_id}.pdf",
                page=1,
                clause="Cl. 1",
                text=f"Trecho de teste para {concept_id}.",
                method="MULTIMODAL",
                confidence=0.9,
            )
        ]
    )
    return ConceptOccurrence(
        concept_id=concept_id,
        term=f"Termo {concept_id}",
        occurrence_type=OccurrenceType.ADDITIONAL_COVERAGE,
        term_relation=TermRelation.EXACT_MATCH,
        contract_status=status,
        evidence=evidence,
        confidence=Level.HIGH,
        amount=amount,
        limit_basis=limit_basis,  # type: ignore[arg-type]
    )


class FakeExtractor:
    """Returns every weighted concept as contracted, except the ones in `overrides`."""

    def __init__(self, overrides: dict[str, ContractStatus] | None = None) -> None:
        self.overrides = overrides or {}

    async def extract(
        self, content: DocumentContent, knowledge_base: KnowledgeBase
    ) -> ExtractionResult:
        items = [
            occurrence(
                c.id, self.overrides.get(c.id, ContractStatus.CONTRACTED), content.document.id
            )
            for c in knowledge_base.weighted
            if self.overrides.get(c.id) != ContractStatus.NOT_FOUND
        ]
        return ExtractionResult(
            intake=DocumentIntake(
                insurer="Seguradora Teste",
                policy_number="123",
                document_type=content.document.type,
            ),
            occurrences=items,
            model="fake",
            prompt_version="test",
        )


class RecordingExtractor(FakeExtractor):
    """FakeExtractor that keeps what the pipeline handed over for each document."""

    def __init__(self) -> None:
        super().__init__()
        self.contents: list[DocumentContent] = []

    async def extract(
        self, content: DocumentContent, knowledge_base: KnowledgeBase
    ) -> ExtractionResult:
        self.contents.append(content)
        return await super().extract(content, knowledge_base)


class FakeAssessor:
    model_name = "fake-assessor"

    async def assess(self, requests: list[AssessmentRequest]) -> list[ConceptAssessment]:
        def decide(side: dict[str, object]) -> AssessmentDecision:
            contracted = side.get("contract_status") == ContractStatus.CONTRACTED
            return AssessmentDecision(
                base_result=1.0 if contracted else 0.25,
                adjustment_factor=1.0,
                justification=None if contracted else "Somente menção.",
            )

        return [
            ConceptAssessment(
                concept_id=r.concept_id,
                a=decide(r.policy_a),
                b=decide(r.policy_b),
                main_difference="",
            )
            for r in requests
        ]


class FakeSummaryWriter:
    async def write_conclusion(self, summary: ExecutiveSummary, facts: dict[str, object]) -> str:
        return summary.conclusion


class FakePdfReader:
    """Stand-in for pypdf: a body of only comments is a valid stub, other text is garbage."""

    def read(self, data: bytes) -> PdfText:
        if b"/Encrypt" in data:
            return PdfText(page_count=0, page_texts={}, encrypted=True)
        body = data.split(b"\n", 1)[1] if b"\n" in data else b""
        if any(line.strip() and not line.startswith(b"%") for line in body.splitlines()):
            return PdfText(page_count=0, page_texts={}, unreadable=True)
        return PdfText(page_count=2, page_texts={})


def build_test_app(tmp_path: Path, extractor: FakeExtractor | None = None) -> FastAPI:
    settings = Settings(_env_file=None, persistence_backend="memory")
    bus = QueuedEventBus(workers=1)
    policies = InMemoryPolicyRepository()
    catalog = JsonConceptCatalog()
    policy_service = PolicyService(
        repository=policies,
        storage=LocalBlobStorage(tmp_path),
        catalog=catalog,
        extractor=extractor or FakeExtractor(),
        pdf_reader=FakePdfReader(),
        docx_reader=PythonDocxTextReader(),
        event_bus=bus,
        limits=UploadLimits(max_file_bytes=1024 * 1024, max_files=5, min_evidence_confidence=0.7),
    )
    comparison_service = ComparisonService(
        comparisons=InMemoryComparisonRepository(),
        policies=policies,
        catalog=catalog,
        assessor=FakeAssessor(),
        summary_writer=FakeSummaryWriter(),
        event_bus=bus,
        parameters=ScoringParameters(),
    )
    bus.subscribe(POLICY_UPLOADED, policy_service.handle_policy_uploaded)
    bus.subscribe(COMPARISON_REQUESTED, comparison_service.handle_comparison_requested)
    services = ApiServices(
        policies=policy_service,
        comparisons=comparison_service,
        concepts=ConceptService(catalog, policies),
    )
    return create_app(settings, bus, services)


PDF_BYTES = b"%PDF-1.4\n% documento de teste\n"
POLICY_TYPE = DocumentType.POLICY


def make_docx(*, with_table: bool = True, header: str | None = "Cabeçalho da Apólice") -> bytes:
    """Build a small policy DOCX in memory (headings, list, table, header, page break)."""

    document = Document()
    if header:
        document.sections[0].header.paragraphs[0].text = header
    document.add_heading("Condições Gerais", level=1)
    document.add_paragraph("1.1 A seguradora cobre perdas de administradores.")
    document.add_paragraph("Custos de defesa", style="List Bullet")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Cobertura"
        table.cell(0, 1).text = "Limite"
        table.cell(1, 0).text = "LMG"
        table.cell(1, 1).text = "R$ 10.000.000"
    document.add_page_break()
    document.add_paragraph("2.1 Exclusões: fraude comprovada.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
